"""The sheet a school hands out, and the two exports that were handing out everyone's.

Reported: "Where is account distribution for students and teacher login distribution?
Maybe here: /admin-sekolah/students and /admin-sekolah/teachers for download login."
Measured before this file existed, and there were two different answers.

**The login a school cannot obtain.** Nothing in this app stores a password it can
read: the credential is a hash in Supabase Auth, and every path that creates an
account generates one and drops it — `_import_students` keeps only
`results["students"] += 1`, and both bulk-reset routes *do* return a password per row
while their pages print "N passwords direset" and discard the rest. So a school could
import two hundred pupils and hold no password for a single one. `/admin-sekolah/students`
and `/admin-sekolah/teachers` therefore gain a **download that issues cards**: it
generates the password, writes it to the account, and carries it into the file in the
same request — the only moment it can exist.

**The login a school could obtain for the whole platform.** `/admin/students/export`
and `/admin/teachers/export` read `profiles` with no school filter, behind
`@admin_required` — which admits `admin_sekolah`. Measured live: the SMP demo admin
downloaded **723 pupils and 73 teachers from every school on the box**, from a link
the admin area itself renders.

What is asserted here, and why each assertion is there:

* the **identity** on a card is the column the login route matches on, not the email —
  `/auth/login-user` finds a pupil by `students.nisn` and a teacher by
  `teachers.employee_id`, so a card printing only the email would look right and not
  work;
* an id from another school is **refused**, not skipped: a sheet quietly short of
  three pupils is one the school hands out and then answers for;
* a reset that failed is **carried in the sheet**, not raised, so one bad account does
  not cost the operator the other ninety-seven passwords it just set;
* the header states that nothing can read a password back, because a school that
  believes it can will file the sheet and delete the file;
* both roster exports scope by the caller's school and leave a super admin's view
  alone.
"""
from __future__ import annotations

import contextlib
import io
import json
import re
from pathlib import Path
from types import SimpleNamespace

import pytest

from tests.conftest import app_instance
from app.utils import auth as auth_utils

ROOT = Path(__file__).resolve().parents[2]
ADMIN_ROUTES = ROOT / "app" / "routes" / "admin_sekolah.py"
LEGACY_ADMIN = ROOT / "app" / "routes" / "admin.py"
AUTH_ROUTES = ROOT / "app" / "routes" / "auth.py"
STUDENTS_PAGE = (ROOT / "app" / "templates" / "admin_sekolah"
                 / "students.html").read_text(encoding="utf-8")
TEACHERS_PAGE = (ROOT / "app" / "templates" / "admin_sekolah"
                 / "teachers.html").read_text(encoding="utf-8")

from app.services import login_cards as lc  # noqa: E402


# ── a supabase stand-in that records writes ─────────────────────────────────

class _Res:
    def __init__(self, data, count=0):
        self.data = data
        self.count = count


class _Query:
    def __init__(self, fake, table):
        self.fake = fake
        self.table = table
        self._eq = []
        self._in = None
        self._update = None

    def select(self, *cols, **kw):
        return self

    def eq(self, col, val):
        self._eq.append((col, val))
        return self

    def in_(self, col, vals):
        self._in = (col, [str(v) for v in vals])
        return self

    def maybe_single(self):
        return self

    def single(self):
        return self

    def update(self, payload):
        self._update = dict(payload)
        return self

    def execute(self):
        rows = [dict(r) for r in self.fake.tables.get(self.table, [])]
        for col, val in self._eq:
            rows = [r for r in rows if r.get(col) == val]
        if self._in:
            col, vals = self._in
            rows = [r for r in rows if str(r.get(col)) in vals]
        if getattr(self, "_update", None) is not None:
            # Issuing a card also marks the account (`must_change_password`), so the
            # stand-in has to accept a write or that step is reported as a failure on
            # every row and the suite measures a fault it invented.
            self.fake.writes.append((self.table, self._update, list(self._eq)))
            return _Res(rows)
        if self._in or self._eq:
            self.fake.reads.append((self.table, list(self._eq), self._in))
        return _Res(rows)


class _Admin:
    def __init__(self, fail_on=()):
        self.fail_on = set(fail_on)
        self.calls = []

    def update_user_by_id(self, uid, payload):
        self.calls.append((uid, payload))
        if uid in self.fail_on:
            raise RuntimeError("auth exploded")
        return type("U", (), {"user": {"id": uid}})()


class _FakeSupabase:
    def __init__(self, tables=None, fail_on=()):
        self.tables = tables or {}
        self.reads = []
        self.writes = []
        self.auth = type("A", (), {"admin": _Admin(fail_on)})()
        self.admin = self.auth.admin

    def table(self, name):
        return _Query(self, name)


SCHOOL = "school-1"

STUDENTS = [
    {"id": "u-1", "nisn": "1001", "school_id": SCHOOL,
     "profiles": {"full_name": "Ana"}, "classes": {"name": "7A"}},
    {"id": "u-2", "nisn": "1002", "school_id": SCHOOL,
     "profiles": {"full_name": "Budi"}, "classes": {"name": "7B"}},
]
TEACHERS = [
    {"id": "t-1", "employee_id": "NIP-1", "school_id": SCHOOL,
     "profiles": {"full_name": "Bu Citra"}, "subjects": {"name": "Fisika"}},
]
OTHER_SCHOOL = [
    {"id": "x-1", "nisn": "9999", "school_id": "school-2",
     "profiles": {"full_name": "Bukan Milik Kita"}, "classes": {"name": "8A"}},
]


def _fake(kind="student", fail_on=()):
    table = "students" if kind == "student" else "teachers"
    return _FakeSupabase({table: (STUDENTS if kind == "student" else TEACHERS)
                          + (OTHER_SCHOOL if kind == "student" else [])},
                         fail_on=fail_on)


def _cards(fake, ids, kind="student", **kw):
    """The route's own composition: read, then write — never the other way round."""
    found = lc.collect(fake, SCHOOL, ids, kind, **kw)
    return lc.set_passwords(fake, found["rows"])


# ── 1. the identity on a card is the one the login page matches ─────────────

class TestWhatALoginCardCarries:
    def test_the_identity_is_the_column_the_login_route_matches(self):
        """Not the email. `/auth/login-user` looks a pupil up by `students.nisn` and a
        teacher by `teachers.employee_id`, so the card must print that column — and
        this reads both sides so the two cannot drift apart."""
        assert lc.KINDS["student"]["identity_column"] == "nisn"
        assert lc.KINDS["teacher"]["identity_column"] == "employee_id"
        auth_source = AUTH_ROUTES.read_text(encoding="utf-8")
        assert '"nisn", "eq", login_input' in auth_source, (
            "the login route stopped matching pupils by NISN")
        assert '.eq("employee_id", login_input)' in auth_source, (
            "the login route stopped matching teachers by employee_id")

    def test_a_student_card_carries_name_identity_class_and_email(self):
        fake = _fake()
        result = lc.collect(fake, SCHOOL, ["u-1"], "student",
                            emails={"u-1": "ana@scan-grade.app"})
        card = result["rows"][0]
        assert card["name"] == "Ana"
        assert card["identity"] == "1001"
        assert card["group"] == "7A"
        assert card["email"] == "ana@scan-grade.app"

    def test_a_teacher_card_carries_the_nip_and_the_subject(self):
        fake = _fake("teacher")
        result = lc.collect(fake, SCHOOL, ["t-1"], "teacher")
        card = result["rows"][0]
        assert card["identity"] == "NIP-1", card
        assert card["group"] == "Fisika", card

# ── 1b. the email on a card, which is how the account is actually signed in ─

class _PagedAdmin:
    """The GoTrue listing as the server really serves it: a page at a time.

    `server_cap` is the point — asking for 1000 does not promise 1000 back, so
    "read the first page and assume it is everything" has to fail here the way it
    fails in production.
    """

    def __init__(self, users, *, server_cap=50):
        self._users = list(users)
        self._cap = server_cap
        self.pages_asked = []

    def list_users(self, page=1, per_page=None):
        per = min(per_page or 50, self._cap)
        self.pages_asked.append((page, per))
        start = (page - 1) * per
        return self._users[start:start + per]


def test_the_email_map_covers_accounts_past_the_first_page(monkeypatch):
    """The sheet is the only place a printed password exists, and the address beside it
    is the other half of the credential: `/auth/login-user` matches on it.

    `_get_email_map` read `auth.admin.list_users()` once, and that returns a page of 50 —
    so on a school of hundreds every card past the first fifty printed an empty Email
    column, which is the one thing the card exists to hand over.
    """
    from app.routes import admin_sekolah as route

    users = [SimpleNamespace(id=f"u-{i}", email=f"guru{i}@scan-grade.app")
             for i in range(120)]
    admin = _PagedAdmin(users, server_cap=50)
    monkeypatch.setattr(auth_utils, "get_auth_admin", lambda: admin)
    monkeypatch.setattr(route, "_email_cache", {"data": {}, "ts": 0.0})

    found = route._get_email_map(None)

    assert found.get("u-119") == "guru119@scan-grade.app", (
        "the map holds only the first page, so every card past it prints no email "
        "and the pupil has no address to sign in with")
    assert len(found) == 120, len(found)


def test_an_official_card_uses_the_auth_address_when_the_mirror_is_empty():
    """Every account on this project predates migration 040's mirror column, so
    `profiles.email` is empty for all of them today — and read alone, it printed a
    blank login address for a head teacher, whose only identity at `/auth/login-user`
    *is* the email.
    """
    fake = _FakeSupabase({"profiles": [
        {"id": "p-1", "full_name": "Kepala Sekolah", "role": "principal",
         "school_id": SCHOOL, "email": None}]})
    result = lc.collect(fake, SCHOOL, ["p-1"], "official",
                        emails={"p-1": "kepsek@scan-grade.app"})
    card = result["rows"][0]
    assert card["identity"] == "kepsek@scan-grade.app", card
    assert card["email"] == "kepsek@scan-grade.app", card


def test_an_official_card_falls_back_to_the_mirror_when_auth_cannot_answer():
    """The mirror is the fallback that keeps a card printable when the auth listing is
    refused — the same reason the column exists.
    """
    fake = _FakeSupabase({"profiles": [
        {"id": "p-2", "full_name": "Wakil Kepala", "role": "vice_principal",
         "school_id": SCHOOL, "email": "wakil@scan-grade.app"}]})
    result = lc.collect(fake, SCHOOL, ["p-2"], "official", emails={})
    assert result["rows"][0]["email"] == "wakil@scan-grade.app", result["rows"][0]


def test_a_student_card_falls_back_to_the_mirror_too():
    """The same two sources, the same order, for the other two kinds — otherwise the
    fallback exists for one role and is forgotten for three.
    """
    fake = _FakeSupabase({"students": [
        {"id": "u-9", "nisn": "1009", "school_id": SCHOOL,
         "profiles": {"full_name": "Citra", "email": "citra@scan-grade.app"},
         "classes": {"name": "7A"}}]})
    result = lc.collect(fake, SCHOOL, ["u-9"], "student", emails={})
    assert result["rows"][0]["email"] == "citra@scan-grade.app", result["rows"][0]


class TestTheEmbeds:
    """What each kind's card joins to, and why the wrong one is not merely empty."""

    def test_the_group_embed_is_the_one_each_kind_has(self):
        """A pupil has a class and a teacher has a subject; asking for the wrong embed
        is a PostgREST error, not an empty column."""
        assert lc.KINDS["student"]["group_embed"] == "classes(name)"
        assert lc.KINDS["teacher"]["group_embed"] == "subjects(name)"


# ── 2. scoping: refused, not skipped ───────────────────────────────────────

class TestWhoseAccountsTheseAre:
    def test_an_id_from_another_school_is_missing_rather_than_read(self):
        fake = _fake()
        result = lc.collect(fake, SCHOOL, ["u-1", "x-1"], "student")
        assert [r["id"] for r in result["rows"]] == ["u-1"]
        assert result["missing"] == ["x-1"], (
            "the other school's pupil was silently dropped instead of being named")

    def test_the_read_is_scoped_by_school(self):
        fake = _fake()
        lc.collect(fake, SCHOOL, ["u-1"], "student")
        table, eqs, _in = fake.reads[-1]
        assert (("school_id", SCHOOL) in eqs), eqs

    def test_no_ids_is_no_query_and_no_rows(self):
        fake = _fake()
        assert lc.collect(fake, SCHOOL, [], "student") == {"rows": [], "missing": []}
        assert fake.reads == [], "an empty request still went to the database"


# ── 3. issuing: the password exists only if this writes it ─────────────────

class TestIssuing:
    def test_every_card_gets_a_password_that_was_written_to_the_account(self):
        fake = _fake()
        cards = _cards(fake, ["u-1", "u-2"])
        printed = [r["password"] for r in cards]
        written = [payload["password"] for _uid, payload in fake.admin.calls]
        assert written == printed, (
            "the sheet and the accounts disagree about the password")
        assert all(len(p) >= 12 for p in printed), printed
        assert len(set(printed)) == 2, "two accounts were given the same password"

    def test_two_runs_do_not_reuse_a_password(self):
        """A card that repeats yesterday's password is a card that opens nothing."""
        first = _cards(_fake(), ["u-1"])[0]["password"]
        second = _cards(_fake(), ["u-1"])[0]["password"]
        assert first != second

    def test_a_reset_that_failed_is_carried_in_the_sheet_not_raised(self):
        """One bad account must not cost the operator the ninety-seven passwords that
        were just set."""
        rows = {r["id"]: r for r in _cards(_fake(fail_on=("u-2",)), ["u-1", "u-2"])}
        assert rows["u-1"]["password"] and not rows["u-1"]["error"]
        assert rows["u-2"]["error"] and not rows["u-2"]["password"], rows["u-2"]

    def test_an_unscoped_id_is_never_issued_a_password(self):
        """The ordering rule, at the service's own level: what is read is what is
        written, so an id the school does not own cannot reach `set_passwords`."""
        fake = _fake()
        found = lc.collect(fake, SCHOOL, ["x-1"], "student")
        lc.set_passwords(fake, found["rows"])
        assert found["missing"] == ["x-1"]
        assert fake.admin.calls == [], (
            "a password was set for an account in another school")


# ── 4. the file ────────────────────────────────────────────────────────────

class TestTheSheet:
    def _sheet(self, kind="student"):
        return _cards(_fake(kind), ["u-1", "u-2"], kind)

    def test_the_csv_has_one_row_per_account_and_every_column(self):
        cards = self._sheet()
        payload = lc.to_csv(cards, "student", lc.meta_for("SMP 1", 2, 2, 0))
        assert payload.startswith(b"\xef\xbb\xbf"), "no BOM — Excel reads mojibake"
        text = payload.decode("utf-8-sig")
        assert "Ana" in text and "1001" in text and cards[0]["password"] in text
        assert "7A" in text and "Budi" in text

    def test_the_xlsx_has_one_row_per_account_and_the_same_columns(self):
        from openpyxl import load_workbook
        cards = self._sheet()
        payload = lc.to_xlsx(cards, "student", lc.meta_for("SMP 1", 2, 2, 0))
        ws = load_workbook(io.BytesIO(payload)).active
        values = list(ws.iter_rows(values_only=True))
        flat = [v for row in values for v in row if v is not None]
        assert "Ana" in flat and "1001" in flat and cards[0]["password"] in flat
        header = [row for row in values
                  if row and row[0] == lc.COLUMNS[0][1][0]][0]
        assert [c for c in header if c] == [label[0] for _key, label in lc.COLUMNS]

    def test_the_two_formats_carry_the_same_rows(self):
        from openpyxl import load_workbook
        cards = self._sheet()
        meta = lc.meta_for("SMP 1", 2, 2, 0)
        csv_rows = [r for r in lc.to_csv(cards, "student", meta)
                    .decode("utf-8-sig").splitlines()
                    if r.startswith(("1,", "2,"))]
        ws = load_workbook(io.BytesIO(lc.to_xlsx(cards, "student", meta))).active
        xlsx_rows = [r for r in ws.iter_rows(values_only=True)
                     if r and r[0] in (1, 2)]
        assert len(csv_rows) == len(xlsx_rows) == 2
        for csv_row, xlsx_row in zip(csv_rows, xlsx_rows):
            assert csv_row.split(",")[1] == xlsx_row[1], (csv_row, xlsx_row)

    def test_the_header_says_the_password_cannot_be_read_back(self):
        """A school that believes it can will file the sheet and delete the file.

        Each language is checked **on its own**: an `or` across the two lets the
        English sheet drop the sentence while the Indonesian one still carries it,
        which is exactly what the mutation harness found on its first run.
        """
        for lang, phrase in (("id", "tidak ada halaman"),
                             ("en", "nothing on the site")):
            meta = lc.meta_for("SMP 1", 2, 2, 0, lang=lang)
            joined = " ".join(f"{k} {v}" for k, v in meta.items()).lower()
            assert "password" in joined, (lang, meta)
            assert phrase in joined, (lang, meta)
            assert "login" in joined, (lang, meta)

    def test_the_header_counts_what_was_requested_issued_and_failed(self):
        meta = lc.meta_for("SMP 1", 5, 4, 1)
        joined = " ".join(str(v) for v in meta.values())
        assert "5" in joined and "4" in joined and "1" in joined

    def test_an_unknown_format_is_xlsx_rather_than_an_error(self):
        """The format comes from a hidden field. A 400 for a typo would cost the
        operator the passwords that were just written."""
        cards = self._sheet()
        payload, mimetype, name = lc.render(cards, "student",
                                            lc.meta_for("SMP 1", 2, 2, 0), "pdf-typo")
        assert name.endswith(".xlsx") and payload[:2] == b"PK"
        assert mimetype == lc.MIMETYPES["xlsx"]

    def test_csv_is_served_as_csv_when_asked(self):
        cards = self._sheet()
        payload, mimetype, name = lc.render(cards, "student",
                                            lc.meta_for("SMP 1", 2, 2, 0), "csv")
        assert name.endswith(".csv") and "text/csv" in mimetype
        assert not payload.startswith(b"PK")

    def test_the_filename_is_ascii_and_says_which_role(self):
        student = lc.filename("student")
        teacher = lc.filename("teacher")
        assert student.startswith("kartu-login-murid-")
        assert teacher.startswith("kartu-login-guru-")
        assert student.isascii() and teacher.isascii()


# ── 5. the routes ──────────────────────────────────────────────────────────

@contextlib.contextmanager
def _post(path, data=None):
    from flask import g
    with app_instance().test_request_context(path, method="POST", data=data or {}):
        g.user_id = "admin-1"
        g.user_role = "admin_sekolah"
        g.user_school_id = SCHOOL
        yield


def _route_modules(monkeypatch, fake, school=SCHOOL):
    from app.routes import admin_sekolah as mod
    monkeypatch.setattr(mod, "get_supabase", lambda: fake)
    monkeypatch.setattr(mod, "_school_id", lambda: school)
    monkeypatch.setattr(mod, "_get_email_map", lambda _s: {})
    monkeypatch.setattr(mod, "log_activity", lambda *a, **k: None)
    return mod


def test_every_card_route_is_registered_and_goes_through_the_one_body():
    """Three doors now, and they must share one body.

    Pupils and teachers were the first two; principal and vice_principal were added
    with the officials work. The count is exact on purpose: a route that grew its own
    implementation would skip the school-scope check, and that check lives only in
    `_login_cards`.
    """
    source = ADMIN_ROUTES.read_text(encoding="utf-8")
    assert '"/students/login-cards"' in source
    assert '"/teachers/login-cards"' in source
    assert '"/officials/login-cards"' in source
    assert source.count("return _login_cards(") == 3, (
        "one of the three card routes stopped sharing the body — the school check "
        "lives there")


def test_a_download_with_no_selection_is_refused_in_words(monkeypatch):
    fake = _fake()
    mod = _route_modules(monkeypatch, fake)
    with _post("/admin-sekolah/students/login-cards", {"user_ids": "[]"}):
        resp = mod._login_cards("student")
        from flask import get_flashed_messages
        flashed = get_flashed_messages(with_categories=True)
    assert resp.status_code == 302, resp.status_code
    assert flashed and flashed[0][1], "nothing said why nothing downloaded"
    assert fake.admin.calls == [], "it reset passwords for an empty selection"


def test_a_download_naming_another_schools_account_is_refused(monkeypatch):
    fake = _fake()
    mod = _route_modules(monkeypatch, fake)
    with _post("/admin-sekolah/students/login-cards",
               {"user_ids": json.dumps(["u-1", "x-1"])}):
        resp = mod._login_cards("student")
        from flask import get_flashed_messages
        flashed = get_flashed_messages(with_categories=True)
    assert resp.status_code == 302, resp.status_code
    assert flashed and flashed[0][0] == "error", flashed
    assert fake.admin.calls == [], (
        "no password may be set while a refusal is being reported")


def test_a_download_returns_a_file_named_for_its_role(monkeypatch):
    fake = _fake()
    mod = _route_modules(monkeypatch, fake)
    with _post("/admin-sekolah/students/login-cards",
               {"user_ids": json.dumps(["u-1", "u-2"]), "fmt": "csv"}):
        resp = mod._login_cards("student")
    assert resp.status_code == 200, resp.status_code
    assert "attachment" in resp.headers["Content-Disposition"]
    assert "kartu-login-murid" in resp.headers["Content-Disposition"]
    # `send_file` streams, so the body has to be collected rather than read off
    # `.data` — Werkzeug refuses implicit conversion in passthrough mode.
    body = b"".join(resp.iter_encoded())
    assert body.startswith(b"\xef\xbb\xbf")
    assert len(fake.admin.calls) == 2, "the file was built without setting the passwords"


# ── 6. the two rosters that covered every school ───────────────────────────

class _ScopeQuery:
    def __init__(self):
        self.eqs = []

    def select(self, *a, **k):
        return self

    def eq(self, col, val):
        self.eqs.append((col, val))
        return self

    def order(self, *a, **k):
        return self


class _ScopeSupabase:
    def __init__(self):
        self.queries = []

    def table(self, name):
        q = _ScopeQuery()
        self.queries.append((name, q))
        return q


@contextlib.contextmanager
def _as(role, school=SCHOOL):
    from flask import g
    with app_instance().test_request_context("/admin/students/export"):
        g.user_role = role
        g.user_school_id = school
        g.user_id = "u-1"
        yield


def test_a_school_admin_export_covers_only_their_school():
    from app.routes import admin as mod
    supabase = _ScopeSupabase()
    with _as("admin_sekolah"):
        mod._export_query(supabase, "profiles", "id, full_name", "murid")
    _table, query = supabase.queries[-1]
    assert ("school_id", SCHOOL) in query.eqs, (
        "the roster export is unscoped again — a school admin downloads every "
        "school's pupils")


def test_a_super_admin_export_is_not_narrowed():
    from app.routes import admin as mod
    supabase = _ScopeSupabase()
    with _as("super_admin"):
        mod._export_query(supabase, "profiles", "id, full_name", "murid")
    _table, query = supabase.queries[-1]
    assert query.eqs == [("role", "murid")], query.eqs


def test_both_legacy_exports_go_through_the_scoped_query():
    source = LEGACY_ADMIN.read_text(encoding="utf-8")
    assert source.count("_export_query(") == 3, (
        "one definition and two call sites; an export that builds its own query is "
        "an export that can forget the school again")


# ── 7. the buttons exist, and say what they do ─────────────────────────────

class TestThePagesOfferIt:
    def test_the_students_page_posts_the_selection_to_the_download(self):
        assert 'action="/admin-sekolah/students/login-cards"' in STUDENTS_PAGE
        assert 'id="student-login-ids"' in STUDENTS_PAGE
        assert "JSON.stringify(selected)" in STUDENTS_PAGE, (
            "the hidden ids field is never filled, so the download would always be "
            "refused as an empty selection")

    def test_the_teachers_page_posts_the_selection_to_the_download(self):
        assert 'action="/admin-sekolah/teachers/login-cards"' in TEACHERS_PAGE
        assert 'id="teacher-login-ids"' in TEACHERS_PAGE
        assert "JSON.stringify(selected)" in TEACHERS_PAGE
        assert "selected: []" in TEACHERS_PAGE, (
            "the teachers table has no selection for the download to use")

    def test_both_buttons_warn_that_the_passwords_change(self):
        for page in (STUDENTS_PAGE, TEACHERS_PAGE):
            m = re.search(r"@submit=\"[^\"]*confirm\(([^\"]+)\)", page)
            assert m, "the download has no confirmation"
            assert "password baru" in m.group(1), m.group(1)

    def test_the_new_label_is_on_both_pages(self):
        for page in (STUDENTS_PAGE, TEACHERS_PAGE):
            assert "Unduh kartu login" in page, page[:80]

    def test_the_english_half_is_only_where_a_reader_can_reach_it(self):
        """A page that pins its language cannot render an English half.

        `base.html`'s `lang` getter returns `data-content-lang` before the reader's
        choice, so a `t()` on a pinned page — `admin_sekolah/teachers.html` pins
        itself, the pupils' page does not — is copy no one can reach. The i18n gate
        agrees: a newly-pinned page carrying pairs is a page that *stopped*
        switching, and it refuses the release. So the pupils' page carries the
        pair and the teachers' page carries the pin instead.
        """
        assert "Download login cards" in STUDENTS_PAGE
        assert "content_lang" not in STUDENTS_PAGE, (
            "the pupils' page pinned itself after all, which freezes the pair just added")
        assert "{% set content_lang = 'id' %}" in TEACHERS_PAGE, (
            "the teachers' page lost its pin, so its Indonesian-only labels are "
            "now an English half nobody wrote")
        assert "t('Unduh kartu login'" not in TEACHERS_PAGE
