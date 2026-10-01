"""Import/upload for teachers, principals and vice principals.

Measured before this file existed: a teacher sheet whose employee number was not
pure digits (`GT-001`, `1987.0101`) fell through every positional branch in
``_import_teachers`` and was read as *"email = the subject column"* — so the row
was handed a subject name as an email address and the whole sheet failed, while
the student sheet (numeric NISN) imported fine. And a school official had no
importer at all: the only way to add a head teacher was the Officials page, one
form submit each.

The guards here hold the header-driven reading and the officials importer, and the
dispatch that routes a "Kepala Sekolah" / "Wakil Kepala Sekolah" / "Pejabat" sheet
to it.
"""
from __future__ import annotations

from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app.routes import admin_sekolah as adm


class _Query:
    def __init__(self, rows):
        self.rows = rows

    def select(self, *a, **k):
        return self

    def eq(self, col, val):
        return self

    def ilike(self, *a, **k):
        return self

    def order(self, *a, **k):
        return self

    def limit(self, *a, **k):
        return self

    def maybe_single(self):
        return self

    def execute(self):
        return SimpleNamespace(data=self.rows, count=len(self.rows))


class _Sb:
    def table(self, name):
        return _Query([])


def _sheet(title, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for row in rows:
        ws.append(row)
    return ws


@pytest.fixture
def captured(monkeypatch):
    """Record what the importers hand the account creators."""
    calls = {"teachers": [], "officials": []}
    monkeypatch.setattr(adm, "create_teacher_account",
                        lambda supabase, **kw: (calls["teachers"].append(kw), "uid")[1])
    from app.services import school_officials
    monkeypatch.setattr(school_officials, "create_official",
                        lambda supabase, **kw: (calls["officials"].append(kw), "uid")[1])
    monkeypatch.setattr(adm, "_get_email_domain", lambda sid: "sekolah.id")
    return calls


class TestTeacherSheetReadByHeader:
    def test_an_alphanumeric_nip_keeps_email_and_subject_in_their_columns(self, captured):
        ws = _sheet("Guru", [
            ["NIP", "Nama Lengkap", "Email", "Mata Pelajaran", "No. HP"],
            ["GT-001", "Budi Santoso", "budi@sekolah.id", "Matematika", "0812"],
        ])
        adm._import_teachers(ws, "S1", _Sb(), {"errors": []})

        call = captured["teachers"][0]
        assert call["employee_id"] == "GT-001"
        assert call["full_name"] == "Budi Santoso"
        assert call["email"] == "budi@sekolah.id", (
            "the subject column was read as the email — the alphanumeric-NIP bug")

    def test_a_float_nip_is_not_printed_with_a_decimal(self, captured):
        ws = _sheet("Guru", [
            ["NIP", "Nama Lengkap", "Email", "Mata Pelajaran"],
            [19870101.0, "Siti Rahma", "siti@sekolah.id", "Fisika"],
        ])
        adm._import_teachers(ws, "S1", _Sb(), {"errors": []})
        assert captured["teachers"][0]["employee_id"] == "19870101"

    def test_a_sheet_without_a_nip_column_still_imports(self, captured):
        """The header names the person; the employee number is not what makes a
        teacher. Measured live: a Guru sheet with a name/email/subject header but
        no NIP column reported "29 baris terbaca tetapi tidak ada yang dikenali"
        while the officials sheet beside it imported, because this reader demanded
        a NIP *column* before it would read the header at all."""
        ws = _sheet("Guru", [
            ["Nama Lengkap", "Email", "Mata Pelajaran", "No. HP"],
            ["Budi Santoso", "budi@sekolah.id", "Matematika", "0812"],
        ])
        adm._import_teachers(ws, "S1", _Sb(), {"errors": []})
        assert len(captured["teachers"]) == 1, captured["teachers"]
        call = captured["teachers"][0]
        assert call["full_name"] == "Budi Santoso"
        assert call["email"] == "budi@sekolah.id"
        assert not call["employee_id"], (
            "no NIP column means no NIP, not the name or the subject read as one")

    def test_a_header_row_whose_nip_cells_are_blank_still_imports(self, captured):
        ws = _sheet("Guru", [
            ["NIP", "Nama Lengkap", "Email"],
            ["", "Siti Rahma", "siti@sekolah.id"],
        ])
        adm._import_teachers(ws, "S1", _Sb(), {"errors": []})
        assert [c["full_name"] for c in captured["teachers"]] == ["Siti Rahma"]

    def test_a_generated_email_is_used_when_the_column_is_blank(self, captured):
        ws = _sheet("Guru", [
            ["NIP", "Nama Lengkap", "Email"],
            ["GT-002", "Tanpa Email", ""],
        ])
        adm._import_teachers(ws, "S1", _Sb(), {"errors": []})
        assert "@sekolah.id" in captured["teachers"][0]["email"]


class TestOfficialSheets:
    def test_a_dedicated_principal_sheet_fixes_the_role(self, captured):
        ws = _sheet("Kepala Sekolah", [
            ["Nama Lengkap", "Email"],
            ["Drs. Hasan", "hasan@sekolah.id"],
        ])
        results = {"officials": 0, "errors": []}
        adm._import_officials(ws, "S1", _Sb(), results, default_role="principal")
        assert results["officials"] == 1
        assert captured["officials"][0]["role"] == "principal"

    def test_a_combined_pejabat_sheet_reads_the_role_column(self, captured):
        ws = _sheet("Pejabat", [
            ["Jabatan", "Nama Lengkap", "Email", "No. HP"],
            ["Kepala Sekolah", "Drs. Hasan", "hasan@sekolah.id", ""],
            ["Wakil Kepala Sekolah", "Rina Marlina", "rina@sekolah.id", "0813"],
        ])
        results = {"officials": 0, "errors": []}
        adm._import_officials(ws, "S1", _Sb(), results, default_role=None)
        roles = [c["role"] for c in captured["officials"]]
        assert roles == ["principal", "vice_principal"], roles

    def test_a_row_with_an_unknown_role_is_skipped_not_guessed(self, captured):
        ws = _sheet("Pejabat", [
            ["Jabatan", "Nama Lengkap", "Email"],
            ["Bendahara", "Orang Lain", "x@sekolah.id"],
        ])
        results = {"officials": 0, "errors": []}
        adm._import_officials(ws, "S1", _Sb(), results, default_role=None)
        assert results["officials"] == 0
        assert captured["officials"] == []


def _upload(app, monkeypatch, wb, calls):
    """POST a workbook to /admin-sekolah/import with every importer stubbed."""
    import io
    from app.utils import auth as authmod

    monkeypatch.setattr(authmod, "_session_for", lambda token: {
        "user_id": "U1", "email": "a@x", "name": "Admin", "role": "admin_sekolah",
        "school_id": "S1", "status": "active",
    })
    monkeypatch.setattr(adm, "get_supabase", lambda: _Sb())
    monkeypatch.setattr(adm, "_school_id", lambda: "S1")
    monkeypatch.setattr(adm, "log_activity", lambda *a, **k: None)
    monkeypatch.setattr(adm, "_import_students",
                        lambda ws, sid, sb, res: calls.setdefault("students", ws.title))
    monkeypatch.setattr(adm, "_import_teachers",
                        lambda ws, sid, sb, res: calls.setdefault("teachers", ws.title))
    monkeypatch.setattr(adm, "_import_subjects",
                        lambda ws, sid, sb, res: calls.setdefault("subjects", ws.title))
    monkeypatch.setattr(adm, "_import_officials",
                        lambda ws, sid, sb, res, default_role=None:
                        calls.setdefault(default_role or "pejabat", ws.title))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    client = app.test_client()
    client.set_cookie("access_token", "tok")
    with client.session_transaction() as sess:
        sess["_csrf_token"] = "csrf"
    return client.post("/admin-sekolah/import",
                       data={"file": (buf, "data.xlsx")},
                       content_type="multipart/form-data",
                       headers={"X-CSRF-Token": "csrf"},
                       follow_redirects=True)


class TestImportDispatch:
    def test_each_named_sheet_reaches_its_importer(self, app, monkeypatch):
        wb = Workbook()
        wb.active.title = "Murid"
        wb.create_sheet("Guru")
        wb.create_sheet("Kepala Sekolah")
        wb.create_sheet("Wakil Kepala Sekolah")
        wb.create_sheet("Mata Pelajaran")
        calls: dict = {}
        _upload(app, monkeypatch, wb, calls)
        assert calls == {
            "students": "Murid", "teachers": "Guru",
            "principal": "Kepala Sekolah",
            "vice_principal": "Wakil Kepala Sekolah",
            "subjects": "Mata Pelajaran",
        }, calls

    def test_a_combined_pejabat_sheet_does_not_fix_the_role(self, app, monkeypatch):
        wb = Workbook()
        wb.active.title = "Pejabat"
        calls: dict = {}
        _upload(app, monkeypatch, wb, calls)
        assert calls == {"pejabat": "Pejabat"}, (
            "a combined officials sheet must let its Role column decide")

    def test_a_copied_sheet_name_still_reaches_its_importer(self, app, monkeypatch):
        # Copying a tab in Excel names it "Guru (2)"; matching the raw name made
        # that a silent zero. Spaces, suffixes and punctuation are not the data.
        wb = Workbook()
        wb.active.title = "Murid "
        wb.create_sheet("Guru (2)")
        wb.create_sheet("MATA PELAJARAN")
        calls: dict = {}
        _upload(app, monkeypatch, wb, calls)
        assert calls == {
            "students": "Murid ", "teachers": "Guru (2)", "subjects": "MATA PELAJARAN",
        }, calls

    def test_a_workbook_with_no_recognisable_sheet_says_so(self, app, monkeypatch):
        # The live complaint: "Impor selesai: 0 murid, 0 guru, 0 mapel. 0 error."
        # Reaching no importer must be reported, not read as a clean run.
        wb = Workbook()
        wb.active.title = "Sheet1"
        wb.create_sheet("Daftar")
        calls: dict = {}
        resp = _upload(app, monkeypatch, wb, calls)
        assert calls == {}
        assert resp.status_code in (302, 200)
        body = _flash_of(resp)
        assert "dikenali" in body.lower(), body
        assert "Sheet1" in body, body

    def test_an_extra_unrecognised_sheet_is_named_not_silently_dropped(self, app, monkeypatch):
        wb = Workbook()
        wb.active.title = "Guru"
        wb.create_sheet("Catatan")
        calls: dict = {}
        resp = _upload(app, monkeypatch, wb, calls)
        assert calls == {"teachers": "Guru"}
        assert "Catatan" in _flash_of(resp)


class TestSilentSkipReported:
    def test_a_teacher_sheet_with_rows_but_no_recognised_header_is_an_error(self, captured):
        ws = _sheet("Guru", [
            ["Kolom A", "Kolom B"],
            ["", "Budi"],
            ["", "Siti"],
        ])
        # two non-blank rows whose positional read leaves the employee number
        # empty, so nothing imports: the sheet must say so rather than report
        # "0 guru, 0 error".
        results = {"teachers": 0, "errors": []}
        adm._import_teachers(ws, "S1", _Sb(), results)
        assert results["teachers"] == 0
        assert any("Guru" in e and "tidak ada yang dikenali" in e.lower()
                   for e in results["errors"]), results["errors"]

    def test_an_official_sheet_with_rows_but_no_role_is_an_error(self, captured):
        ws = _sheet("Pejabat", [
            ["Jabatan", "Nama Lengkap", "Email"],
            ["Bendahara", "Orang Lain", "x@sekolah.id"],
        ])
        results = {"officials": 0, "errors": []}
        adm._import_officials(ws, "S1", _Sb(), results, default_role=None)
        assert results["officials"] == 0
        assert any("Pejabat" in e and "tidak ada yang dikenali" in e.lower()
                   for e in results["errors"]), results["errors"]


def _flash_of(resp):
    """The page body after the import redirect — the flash message lives in it."""
    return resp.get_data(as_text=True)


class TestCombinedTemplate:
    def test_the_full_template_holds_every_named_sheet(self, app, monkeypatch):
        import io
        from openpyxl import load_workbook as _load
        from app.utils import auth as authmod

        monkeypatch.setattr(authmod, "_session_for", lambda token: {
            "user_id": "U1", "email": "a@x", "name": "Admin",
            "role": "admin_sekolah", "school_id": "S1", "status": "active",
        })
        monkeypatch.setattr(adm, "get_supabase", lambda: _Sb())
        monkeypatch.setattr(adm, "_school_id", lambda: "S1")
        monkeypatch.setattr(adm, "_get_email_domain", lambda sid: "sekolah.id")

        client = app.test_client()
        client.set_cookie("access_token", "tok")
        resp = client.get("/admin-sekolah/download-template/semua")
        assert resp.status_code == 200, resp.status_code
        wb = _load(io.BytesIO(resp.data))
        assert wb.sheetnames == ["Murid", "Guru", "Pejabat", "Mata Pelajaran"], wb.sheetnames
        assert wb["Guru"].cell(row=1, column=4).value == "Mata Pelajaran"


class TestHeaderDetection:
    def test_headers_are_matched_ignoring_case_and_punctuation(self):
        ws = _sheet("Guru", [["No. HP", "Nama  Lengkap", "NIP", "MATA PELAJARAN"]])
        columns = adm._header_columns(ws)
        assert columns["phone"] == 0
        assert columns["name"] == 1
        assert columns["nip"] == 2
        assert columns["subject"] == 3
