"""Four frameworks named, and the two the app never published computed honestly.

One paper's answers can be read four ways, and the app only ever did two of them
silently. This file is the vocabulary and the arithmetic behind the other two:

* the **framework catalogue** — every framework says what question it answers,
  what it measures, when a school uses it, what it *cannot* answer and which
  panels it shows, in both languages;
* the **cognitive mix** — HOTS/MOTS/LOTS by marks when the paper carries
  weights and by question count when it does not, with the unlabelled questions
  left visible rather than counted as lower order;
* **criterion-referenced mastery** — tuntas/belum tuntas against the school's own
  KKM, and which questions the class has mastered by that same standard.

Two directions are checked everywhere, because both are defects: a level that is
*guessed* from the question type invents the thing the framework measures, and a
missing KKM read as "everything passes" invents a standard nobody set.
"""
from __future__ import annotations

import re
from pathlib import Path

import pytest
from flask import Flask

from app.services import analysis_frameworks as af
from app.services import item_analysis as ia

ROOT = Path(__file__).resolve().parents[2]
FORM = ROOT / "app" / "templates" / "teacher" / "exam_form.html"
ROUTE = ROOT / "app" / "routes" / "teacher.py"



# ── helpers ──────────────────────────────────────────────────────────────────

def exam(n: int = 4, weights: dict | None = None, key: dict | None = None,
         levels: dict | None = None, kkm: int | None = 70) -> dict:
    row = {
        "id": "exam-1", "title": "Mid Semester 1", "total_questions": n,
        "question_types": {str(i): "mcq" for i in range(n)},
        "question_weights": weights if weights is not None
        else {str(i): 25.0 for i in range(n)},
        "answer_key": key if key is not None else {str(i): "A" for i in range(n)},
    }
    if levels is not None:
        row["question_cognitive"] = levels
    if kkm is not None:
        row["passing_score"] = kkm
    return row


def papers(rows: list[list]) -> list[dict]:
    """`1` right, `0` wrong, `None` blank — the same shape the routes load."""
    out = []
    for j, row in enumerate(rows):
        answers = {str(i): ("A" if value else "B")
                   for i, value in enumerate(row) if value is not None}
        out.append({"id": f"s{j}", "student_id": f"s{j}",
                    "student_name": f"Paper {j}", "answers": answers})
    return out


def analyse(**kwargs) -> ia.Analysis:
    rows = kwargs.pop("rows", [[1, 1, 0, 0], [1, 0, 1, 0], [0, 1, 0, 1],
                               [1, 1, 1, 1], [0, 0, 1, 0], [1, 0, 0, 1]])
    return ia.analyse(exam(**kwargs), papers(rows))


# ── the catalogue ────────────────────────────────────────────────────────────

class TestTheFrameworkCatalogue:
    def test_there_are_four_and_they_are_the_ones_a_school_asks_for(self):
        assert [framework.key for framework in af.FRAMEWORKS] == [
            "ctt", "rasch", "cognitive", "mastery"]
        assert af.DEFAULT == "ctt"

    def test_each_one_says_what_it_answers_in_both_languages(self):
        for framework in af.FRAMEWORKS:
            for field in ("name", "question", "not_for", "when", "reference"):
                pair = getattr(framework, field)
                assert set(pair) == {"id", "en"}, f"{framework.key}.{field}"
                for half in ("id", "en"):
                    assert pair[half].strip(), f"{framework.key}.{field}.{half}"

    def test_each_one_names_what_it_measures_and_what_it_cannot(self):
        """`not_for` is not decoration: it is what stops a reader treating a logit
        as a mark, and it is the reason the four are not interchangeable."""
        for framework in af.FRAMEWORKS:
            assert framework.measures, f"{framework.key} measures nothing"
            assert all(measure.strip() for measure in framework.measures)

    def test_every_panel_a_framework_claims_is_a_panel_that_exists(self):
        for framework in af.FRAMEWORKS:
            assert framework.panels, f"{framework.key} shows no panel"
            for panel in framework.panels:
                assert panel in af.PANELS, f"{framework.key} claims unknown {panel}"

    def test_the_item_table_is_on_every_framework(self):
        """It is the spine: P and D for CTT, b and fit for Rasch, the level for
        the cognitive read, the verdict for mastery. A framework without it would
        be a summary with nothing under it."""
        for framework in af.FRAMEWORKS:
            assert framework.shows("items"), f"{framework.key} hides the item table"

    def test_the_two_rasch_only_blocks_do_not_leak_into_the_other_frameworks(self):
        """Separation and the logit series are statements about a *scale*; showing
        them under CTT would publish a logit to a reader who asked for a
        percentage.

        The item map is not one of them: its two axes are P and D, which is
        exactly what CTT publishes, so a classical reader gets it too. The panel
        docstring names it "difficulty against discrimination" for that reason.
        """
        assert af.BY_KEY["rasch"].shows("separation")
        assert af.BY_KEY["rasch"].shows("difficulty")
        assert af.BY_KEY["ctt"].shows("map"), "the item map is a classical plot"
        for key in ("ctt", "cognitive", "mastery"):
            assert not af.BY_KEY[key].shows("separation"), key
            assert not af.BY_KEY[key].shows("difficulty"), key

    def test_resolve_takes_anything_and_lands_on_the_default(self):
        assert af.resolve("rasch").key == "rasch"
        assert af.resolve("  RASCH ").key == "rasch"
        assert af.resolve("").key == af.DEFAULT
        assert af.resolve(None).key == af.DEFAULT
        assert af.resolve("../../etc/passwd").key == af.DEFAULT

    def test_the_words_of_one_framework_are_one_language(self):
        """A document is already in a language; handing it a pair would put both
        halves in the file."""
        for lang, expected in (("id", "Teori Tes Klasik (CTT)"),
                               ("en", "Classical test theory (CTT)")):
            words = af.BY_KEY["ctt"].words(lang)
            assert words["name"] == expected
            assert "·" in words["measures"]
            assert all(isinstance(value, str) for value in words.values())
        assert af.BY_KEY["ctt"].words("klingon")["name"].startswith("Teori")


# ── the cognitive vocabulary ─────────────────────────────────────────────────

class TestTheCognitiveVocabulary:
    def test_six_levels_in_three_bands(self):
        assert [level.key for level in af.LEVELS] == ["c1", "c2", "c3", "c4", "c5", "c6"]
        assert af.BANDS == ("lots", "mots", "hots")
        assert [af.band_of(level.key) for level in af.LEVELS] == [
            "lots", "lots", "mots", "mots", "hots", "hots"]

    def test_every_band_has_a_name_in_both_languages(self):
        for band in af.BANDS:
            assert set(af.BAND_NAMES[band]) == {"id", "en"}
            for half in ("id", "en"):
                assert af.BAND_NAMES[band][half].strip()

    def test_a_level_carries_its_name_and_the_verb_it_asks_for(self):
        c5 = af.BY_LEVEL["c5"]
        assert c5.band == "hots"
        assert c5.name["en"] == "C5 Evaluating"
        assert "justify" in c5.verb["en"] or "judge" in c5.verb["en"], \
            "the verb is what makes a level checkable against a question"

    def test_an_unset_level_is_none_and_not_a_guess(self):
        """The defect this guards: reading a level off the question type. A
        multiple-choice question is not automatically LOTS, and reporting it as
        such would answer the framework's question with a fabrication."""
        for empty in (None, "", "   ", "c9", "lots", 7):
            assert af.level(empty) is None, empty
            assert af.band_of(empty) is None, empty
        assert af.level("C5").key == "c5"
        assert af.level(" c5 ") is not None

    def test_the_payloads_the_page_and_forms_read(self):
        levels = af.levels_payload()
        assert [item["key"] for item in levels] == [level.key for level in af.LEVELS]
        assert all(set(item["name"]) == {"id", "en"} for item in levels)
        bands = af.bands_payload()
        assert [band["key"] for band in bands] == list(af.BANDS)

    def test_the_catalogue_in_one_language_is_what_a_document_gets(self):
        rows = af.catalogue("en")
        assert len(rows) == 4
        assert {row["key"] for row in rows} == {f.key for f in af.FRAMEWORKS}
        assert all(row["name"] for row in rows)


# ── the cognitive mix ────────────────────────────────────────────────────────

class TestTheCognitiveMix:
    def test_the_mix_counts_marks_when_the_paper_carries_weights(self):
        result = analyse(levels={"0": "c1", "1": "c4", "2": "c5", "3": "c6"},
                         weights={"0": 10.0, "1": 10.0, "2": 30.0, "3": 50.0})
        mix = result.cognitive
        assert mix.basis == "marks"
        assert mix.counts == {"lots": 1, "mots": 1, "hots": 2}
        assert mix.marks == {"lots": 10.0, "mots": 10.0, "hots": 80.0}
        assert mix.share("hots") == 80.0
        assert mix.share("lots") == 10.0

    def test_the_mix_counts_questions_when_the_paper_has_no_weights(self):
        """An exam saved before weights existed has every mark at zero, and a
        share of zero marks is 0.0% for every band — a paper reported as having
        no higher-order thinking at all because nobody wrote a weight down."""
        result = analyse(levels={"0": "c5", "1": "c5", "2": "c1", "3": "c1"},
                         weights={})
        mix = result.cognitive
        assert mix.basis == "questions"
        assert mix.share("hots") == 50.0
        assert mix.share("lots") == 50.0

    def test_unlabelled_questions_are_not_counted_as_lower_order(self):
        result = analyse(levels={"0": "c5"})
        mix = result.cognitive
        assert mix.unset == 3
        assert mix.labelled == 1
        assert mix.counts["lots"] == 0
        assert mix.counts["hots"] == 1
        assert mix.share("hots") == 25.0, \
            "the share is of the paper, so an unlabelled question still counts"

    def test_a_paper_with_no_labels_at_all_says_so(self):
        result = analyse()
        assert result.cognitive.unset == 4
        assert result.cognitive.labelled == 0
        assert all(result.cognitive.counts[band] == 0 for band in af.BANDS)
        assert result.cognitive.share("hots") == 0.0
        assert "levels_partly_set" not in result.notes, \
            "nothing is set, so nothing is partly set either"

    def test_a_partly_labelled_paper_carries_a_note(self):
        assert "levels_partly_set" in analyse(levels={"0": "c5"}).notes
        assert "levels_partly_set" not in analyse(
            levels={str(i): "c5" for i in range(4)}).notes

    def test_the_level_and_band_reach_the_item(self):
        result = analyse(levels={"0": "c6"})
        first, second = result.items[0], result.items[1]
        assert (first.level, first.band) == ("c6", "hots")
        assert (second.level, second.band) == ("", ""), \
            "an unlabelled question stays unlabelled"

    def test_a_garbage_level_does_not_label_an_item(self):
        result = analyse(levels={"0": "C7", "1": "higher", "2": 5})
        assert all(item.band == "" for item in result.items)
        assert result.cognitive.unset == 4


# ── criterion-referenced mastery ─────────────────────────────────────────────

class TestMastery:
    def test_the_mastered_share_of_a_paper(self):
        result = analyse(kkm=60)
        mastery = result.mastery
        assert mastery.kkm == 60
        assert mastery.configured is True
        assert mastery.passed + mastery.failed == result.summary.students
        assert mastery.items_measured == 4, "every question was answered by somebody"

    def test_the_threshold_is_the_kkm_itself(self):
        """A question is mastered when the share of papers earning full credit on
        it reaches the school's own standard — not a number this app invented."""
        assert ia.Mastery(kkm=70).threshold == pytest.approx(0.7)
        assert ia.Mastery(kkm=100).threshold == pytest.approx(1.0)
        assert ia.Mastery(kkm=0).threshold == 0.0

    def test_a_question_is_mastered_by_the_share_who_earned_full_credit(self):
        """Six papers, KKM 70: four right is 66.7% (not mastered), five is 83.3%
        (mastered). Stated as the arithmetic rather than as a flag, because the
        denominator — papers that answered — is the part that goes wrong."""
        rows = [[1, 1, 1, 0], [1, 0, 1, 0], [1, 1, 1, 1], [1, 0, 1, 0],
                [0, 0, 1, 1], [1, 0, 0, 1]]
        result = ia.analyse(exam(4, kkm=70), papers(rows))
        full = [item.full for item in result.items]
        answered = [item.answered for item in result.items]
        assert full == [5, 2, 5, 3], full
        assert answered == [6] * 4
        assert [item.mastered for item in result.items] == [True, False, True, False]
        assert result.mastery.items_mastered == 2
        assert result.mastery.items_measured == 4

    def test_a_question_nobody_answered_has_no_verdict(self):
        rows = [[None, 1, 1, 1], [None, 0, 1, 1], [None, 1, 0, 1]]
        result = ia.analyse(exam(4, kkm=70), papers(rows))
        assert result.items[0].answered == 0
        assert result.items[0].mastered is None, \
            "no data is not a fail — it is a question with no verdict"
        assert result.mastery.items_measured == 3

    def test_an_exam_with_no_kkm_reports_no_verdict_rather_than_everything_passing(self):
        result = analyse(kkm=None)
        mastery = result.mastery
        assert mastery.kkm == 0
        assert mastery.configured is False
        assert mastery.passed == 0 and mastery.failed == 0
        assert mastery.items_mastered == 0
        assert all(item.mastered is None for item in result.items), \
            "a missing standard is not a low bar"
        assert "kkm_missing" in result.notes

    @pytest.mark.parametrize("value,expected", [(0, 0), (None, 0), ("", 0),
                                               ("abc", 0), (-5, 0), (150, 0),
                                               (70, 70), ("75", 75), (72.6, 72)])
    def test_an_unreadable_or_impossible_kkm_is_not_configured(self, value, expected):
        assert ia._kkm(value) == expected

    def test_the_gap_and_the_pass_rate_are_signed_and_stated(self):
        mastery = ia.Mastery(kkm=70, configured=True, passed=3, failed=1, mean=74.5)
        assert mastery.gap == 4.5
        assert mastery.pass_rate == 75.0
        below = ia.Mastery(kkm=70, configured=True, passed=1, failed=3, mean=61.0)
        assert below.gap == -9.0
        assert below.pass_rate == 25.0
        assert ia.Mastery(kkm=70).gap is None
        assert ia.Mastery(kkm=70).pass_rate is None

    def test_the_class_mean_and_the_lowest_paper_come_from_the_scores(self):
        rows = [[1, 1, 1, 1], [0, 0, 0, 0], [1, 0, 1, 0]]
        result = ia.analyse(exam(4, kkm=50), papers(rows))
        scores = [person.score for person in result.people]
        assert result.mastery.mean == round(sum(scores) / len(scores), 2)
        assert result.mastery.lowest == round(min(scores), 2)

    def test_moving_the_kkm_moves_the_verdict_and_nothing_else(self):
        """The framework's own warning, as a test: the answers did not change."""
        strict = analyse(kkm=90)
        gentle = analyse(kkm=30)
        assert strict.mastery.passed <= gentle.mastery.passed
        assert strict.summary.total_mean == gentle.summary.total_mean
        assert [item.pct for item in strict.items] == [item.pct for item in gentle.items]

    def test_a_paper_with_no_submissions_is_not_a_crash(self):
        result = ia.analyse(exam(4, kkm=70), [])
        assert result.mastery.mean is None
        assert result.mastery.passed == 0
        assert result.mastery.items_measured == 0
        assert result.cognitive.unset == 4


# ── the teacher's half: labelling the questions ───────────────────────────

class TestTheBuilderCollectsTheLabels:
    """A framework nobody can feed is a framework that reports nothing.

    The cognitive mix is only as true as the labels behind it, and the labels come
    from one place: the exam builder. Three things have to hold, and each one fails
    silently on its own — the page offers the server's six levels rather than a
    hand-written list, the labels travel to the save route under a name the route
    reads, and the route stores only levels the vocabulary knows.
    """

    def form(self):
        return FORM.read_text(encoding="utf-8")

    def test_the_page_offers_the_servers_own_levels(self):
        text = self.form()
        assert "const SG_LEVELS = {{ cognitive_levels() | tojson | safe }};" in text
        # The select is built by looping that list. A literal "C1 Mengingat" in the
        # template would be a second list of names, and a second list is how the
        # builder offers a level the report cannot group.
        assert "C1 Mengingat" not in text
        # And the loop is *inside* the control bound to the question's own field — a
        # select that lists the levels but writes somewhere else labels nothing.
        select = text.split('<select x-model="q.level"', 1)[1].split("</select>", 1)[0]
        assert 'x-for="lv in SG_LEVELS"' in select
        # The blank first option is the "not labelled yet" choice, and it has to be
        # the one the teacher sees before touching anything.
        assert select.index('<option value=""') < select.index('x-for="lv in SG_LEVELS"')

    def test_the_labels_travel_under_the_name_the_route_reads(self):
        text = self.form()
        assert 'name="question_cognitive"' in text
        assert "getLevelsJson()" in text
        # The builder's key is the hidden field's binding, not the label's — the
        # field is what the form actually posts.
        assert 'id="question-cognitive-input" :value="getLevelsJson()"' in text

    def test_an_unlabelled_question_is_posted_as_nothing(self):
        """Blank must not become a value: absent is "not labelled yet"."""
        text = self.form()
        block = text.split("getLevelsJson() {", 1)[1].split("},", 1)[0]
        assert ".filter(([k,v]) => v)" in block

    def test_a_restored_exam_keeps_its_labels(self):
        """Reopening a labelled paper and saving it must not wipe the kisi-kisi."""
        text = self.form()
        assert "exam.question_cognitive" in text
        assert "level: plevel[i] || ''" in text
        assert "level: ''" in text  # a newly added question starts unlabelled


class TestTheRouteKeepsTheLabelsHonest:

    def labels(self, posted):
        from app.routes.teacher import _cognitive_levels
        app = Flask(__name__)
        with app.test_request_context("/x", method="POST", data=posted):
            return _cognitive_levels()

    def test_a_known_level_is_kept_under_the_question_index(self):
        assert self.labels({"question_cognitive": '{"0": "c4", "2": "c6"}'}) == \
            {"0": "c4", "2": "c6"}

    def test_an_unknown_level_is_dropped_rather_than_stored(self):
        """A level no band knows would count in none of them while looking set."""
        assert self.labels({"question_cognitive": '{"0": "c9", "1": "hots"}'}) == {}

    def test_nonsense_is_an_empty_map_not_an_exception(self):
        assert self.labels({}) == {}
        assert self.labels({"question_cognitive": "not json"}) == {}
        assert self.labels({"question_cognitive": '["c4"]'}) == {}

    def test_a_non_string_level_is_not_a_level(self):
        assert self.labels({"question_cognitive": '{"0": 4}'}) == {}

    def test_the_case_a_teacher_typed_does_not_matter(self):
        assert self.labels({"question_cognitive": '{"0": "C4"}'}) == {"0": "c4"}

    def test_both_save_routes_write_the_column(self):
        """Create and edit are two code paths, and only one used to exist."""
        text = ROUTE.read_text(encoding="utf-8")
        assert text.count('"question_cognitive": _cognitive_levels(),') == 2

    def test_the_edit_form_is_handed_an_object_for_this_column_too(self):
        """Same defect `question_pages` had: a JSON string indexed by a key."""
        from app.routes.teacher import JSON_COLUMNS, _normalise_exam_json
        assert "question_cognitive" in JSON_COLUMNS
        row = {"question_cognitive": '{"0": "c4"}'}
        assert isinstance(_normalise_exam_json(row)["question_cognitive"], dict)
        assert _normalise_exam_json(row)["question_cognitive"] == {"0": "c4"}

    def test_the_column_exists_and_defaults_to_empty(self):
        sql = (ROOT / "supabase" / "migrations"
               / "032_question_cognitive_level.sql").read_text(encoding="utf-8")
        assert re.search(r"ALTER TABLE exams ADD COLUMN IF NOT EXISTS "
                         r"question_cognitive JSONB DEFAULT '\{\}'", sql)
        # Additive and repeatable: nothing is backfilled, because no level of an
        # already-written question is in this database to backfill from.
        assert "UPDATE" not in sql.upper()


class TestTheReportShowsTheChosenFramework:
    """One exam, four answers — and the page renders only the one asked for.

    The panel list is a promise about what a reader sees, and the failure it can
    have is silence: a panel that renders for every framework is not a framework
    choice, and a panel that renders for none is a blank page. Both directions are
    checked on the *rendered* page, not on the template's text.
    """

    EXAM = {
        "id": "exam-1", "title": "Mid Semester 1", "subject": "Fisika",
        "total_questions": 4,
        "question_types": {str(i): "mcq" for i in range(4)},
        "question_weights": {str(i): 25.0 for i in range(4)},
        "answer_key": {str(i): "A" for i in range(4)},
        "question_cognitive": {"0": "c1", "1": "c4", "2": "c6"},
        "passing_score": 70,
    }

    @pytest.fixture(scope="module")
    def analysis(self):
        rows = [[1, 1, 0, 0], [1, 0, 1, 0], [0, 1, 0, 1],
                [1, 1, 1, 1], [0, 0, 1, 0], [1, 0, 0, 1]]
        return ia.analyse(self.EXAM, papers(rows))

    @pytest.fixture(scope="module")
    def app(self):
        from app import create_app
        return create_app("app.config.TestingConfig")

    def render(self, app, analysis, key):
        import contextlib

        from flask import g

        from app.routes.teacher import _chart_payload

        framework = af.resolve(key)
        chart = _chart_payload(analysis, framework=framework)
        with app.test_request_context("/teacher/analysis/exam-1"):
            g.user_id = "tea-1"
            g.user_name = "Guru Uji"
            g.user_email = "guru@example.test"
            g.user_role = "guru"
            g.tz_offset = 7
            g.show = {}
            with contextlib.suppress(Exception):
                g.user_school_id = "sch-1"
            return app.jinja_env.get_template("teacher/analysis.html").render(
                exam=self.EXAM, analysis=analysis, chart=chart, framework=framework)

    def markers(self, html):
        return set(re.findall(r'data-panel="(\w+)"', html))

    @pytest.mark.parametrize("key", [f.key for f in af.FRAMEWORKS])
    def test_a_framework_renders_exactly_its_own_panels(self, app, analysis, key):
        framework = af.resolve(key)
        shown = self.markers(self.render(app, analysis, key))
        assert shown, f"{key} rendered no panel at all"
        for panel in shown:
            assert framework.shows(panel), \
                f"{key} renders the {panel} panel, which it does not claim"
        for panel in af.PANELS:
            if not framework.shows(panel):
                assert panel not in shown, \
                    f"{key} hides {panel} in its promise but renders it"

    def test_the_four_views_really_do_differ(self, app, analysis):
        """Otherwise the selector is decoration: a reader would be choosing
        between four names for one page."""
        seen = {key: self.markers(self.render(app, analysis, key))
                for key in ("ctt", "rasch", "cognitive", "mastery")}
        assert len(set(map(frozenset, seen.values()))) == 4, seen

    def test_the_framework_reaches_the_page_and_its_downloads(self, app, analysis):
        html = self.render(app, analysis, "rasch")
        assert 'data-active="true"' in html or "data-framework-picker" in html
        # The download links carry it, or a teacher reading the Rasch view files a
        # document that names CTT. `chart.` and not `data.`: the payload is a
        # closure variable inside the component, and an Alpine expression can only
        # see a property on it — `data.framework` in an `:href` renders the literal
        # `undefined` into the URL, silently.
        assert html.count("'&framework=' + (chart.framework || '')") == 3
        assert "data.framework || ''" not in html
        # And the same for the selector: an `x-for` over a property-less payload
        # renders nothing, which looks exactly like a page with no choices.
        assert "x-for=\"f in chart.frameworks\"" in html
        assert "x-for=\"f in data.frameworks\"" not in html
        # And a marker that names a *key* has to be bound, or the attribute is the
        # literal text `band.key` on every bar — which reads as correct in the
        # source and as one band in the DOM. Read on the two views that render
        # them, since a hidden panel renders no marker at all.
        from app.services import analysis_frameworks as frameworks
        text = (ROOT / "app" / "templates" / "teacher" / "analysis.html") \
            .read_text(encoding="utf-8")
        # `data-band="band.key"` is a *substring* of the bound form, so the check
        # is on the counts: every occurrence carries its colon.
        assert text.count('data-band="band.key"') == 1
        assert text.count(':data-band="band.key"') == 1
        assert text.count("data-mastery=") == text.count(":data-mastery=") == 1
        assert frameworks.resolve("cognitive").shows("cognitive")
        rendered = self.render(app, analysis, "cognitive")
        assert ':data-band="band.key"' in rendered, \
            "the cognitive panel renders no band marker to hang a key on"
        # One marker in the source, cloned per band in the DOM — so the count is
        # checked in the browser (`preview_evaluate`) and not here, where a `x-for`
        # body is a single occurrence however many bands there are.

    def test_the_pasted_figure_names_the_framework(self, app, analysis):
        """A figure in somebody's report has to say which reading it is; the copy
        board cannot reach the page's header, so the payload carries it and the
        paint routine draws it."""
        html = self.render(app, analysis, "rasch")
        assert "frameworkLine() {" in html, "the pasted figure names no framework"
        assert "text(this.frameworkLine()" in html, \
            "the figure line is defined but never drawn"
        assert "Rasch" in html, "the payload carries no framework name to draw"

    def test_a_reading_with_no_kkm_says_so_instead_of_zeroing(self, app):
        """The mastery framework on a paper whose school never set a KKM."""
        exam = dict(self.EXAM, passing_score=0)
        rows = [[1, 1, 0, 0], [1, 0, 1, 0], [0, 1, 0, 1],
                [1, 1, 1, 1], [0, 0, 1, 0], [1, 0, 0, 1]]
        result = ia.analyse(exam, papers(rows))
        assert result.mastery.configured is False
        html = self.render(app, result, "mastery")
        assert 'data-panel="mastery"' in html
        # The sentence that replaces the verdict, in both languages, on the page
        # rather than only in the file: four cards reading zero would look like a
        # standard everybody failed.
        assert "This exam sets no KKM, so mastery is not computed" in html
        assert "Ujian ini belum menetapkan KKM" in html


def _call_arguments(source: str, opener: str) -> list[str]:
    """The argument text of every call to `opener`, parens matched.

    Written out rather than regexed because a regex that stops at the first `)`
    reads a nested call as the whole argument list — which is how a check like
    this reports success on a call whose framework argument was deleted.
    """
    found = []
    start = source.find(opener)
    while start != -1:
        depth, index = 1, start + len(opener)
        while index < len(source) and depth:
            if source[index] == "(":
                depth += 1
            elif source[index] == ")":
                depth -= 1
            index += 1
        found.append(source[start + len(opener):index - 1])
        start = source.find(opener, index)
    return found


def pdf_text(blob: bytes) -> str:
    """Every word a generated PDF draws, so its content can be asserted.

    Read rather than trusted: the report is built by reportlab and the only
    alternative to reading it is asserting that a function was *called*, which
    stays green while the document prints nothing. `PyMuPDF` is already a
    requirement of this app (it renders the answer sheets), so this adds no
    dependency a deploy would have to install.
    """
    import fitz

    with fitz.open(stream=blob, filetype="pdf") as document:
        return "\n".join(page.get_text() for page in document)


class TestTheDocumentsCarryTheFramework:
    """A filed report has to say *which* reading it is, and print that one.

    The page's promise is per framework, and the three documents are built on the
    server where the framework cannot be guessed from the tab. A CSV named "Item
    Analysis" that carries the HOTS breakdown the classical view never showed is a
    document nobody can date — and a PDF whose header names one framework while
    its tables print another is worse than either.
    """

    EXAM = {
        "id": "exam-1", "title": "Mid Semester 1", "subject": "Fisika",
        "total_questions": 4,
        "question_types": {str(i): "mcq" for i in range(4)},
        "question_weights": {str(i): 25.0 for i in range(4)},
        "answer_key": {str(i): "A" for i in range(4)},
        "question_cognitive": {"0": "c1", "1": "c4", "2": "c6"},
        "passing_score": 70,
    }

    @pytest.fixture(scope="module")
    def analysis(self):
        rows = [[1, 1, 0, 0], [1, 0, 1, 0], [0, 1, 0, 1],
                [1, 1, 1, 1], [0, 0, 1, 0], [1, 0, 0, 1]]
        return ia.analyse(self.EXAM, papers(rows))

    def csv(self, analysis, key, lang="en"):
        from app.services import analysis_report
        return analysis_report.analysis_csv(
            analysis, self.EXAM, lang, framework=af.resolve(key))

    def test_the_csv_names_the_framework_it_was_read_as(self, analysis):
        for key in ("ctt", "rasch", "cognitive", "mastery"):
            text = self.csv(analysis, key)
            assert af.resolve(key).name["en"] in text, key

    def test_the_csv_carries_only_its_own_frameworks_blocks(self, analysis):
        from app.services import analysis_report
        t = analysis_report.labels("en")
        assert t["cognitive"] in self.csv(analysis, "cognitive")
        assert t["cognitive"] not in self.csv(analysis, "ctt")
        assert t["mastery"] in self.csv(analysis, "mastery")
        assert t["mastery"] not in self.csv(analysis, "rasch")

    def test_the_unlabelled_row_is_in_the_mix_not_folded_into_lots(self, analysis):
        """Question 3 has no level: the paper is *partly* labelled, and the file
        has to carry that rather than counting it as lower order."""
        from app.services import analysis_report
        text = self.csv(analysis, "cognitive")
        assert analysis_report.labels("en")["unlabelled"] in text
        assert analysis.cognitive.unset == 1
        assert "LOTS" in text

    def test_the_mastery_block_says_when_no_standard_was_set(self, analysis):
        from app.services import analysis_report
        exam = dict(self.EXAM, passing_score=0)
        rows = [[1, 1, 0, 0], [1, 0, 1, 0], [0, 1, 0, 1],
                [1, 1, 1, 1], [0, 0, 1, 0], [1, 0, 0, 1]]
        result = ia.analyse(exam, papers(rows))
        text = analysis_report.analysis_csv(
            result, exam, "en", framework=af.resolve("mastery"))
        # The sentence is the reason, and no count rows follow it: five rows of
        # zeros would read as a standard nobody met.
        assert "sets no KKM, so mastery is not computed" in text
        labels = analysis_report.labels("en")
        for row in (labels["passed"], labels["failed"], labels["class_mean"],
                    labels["threshold"]):
            assert row not in text, f"an unconfigured KKM still printed {row}"

    def test_the_workbook_keeps_the_blocks_with_their_own_sheets(self, analysis):
        from openpyxl import load_workbook
        from io import BytesIO

        from app.services import analysis_report

        book = analysis_report.analysis_xlsx(
            analysis, self.EXAM, lang="en", framework=af.resolve("cognitive"))
        sheets = load_workbook(BytesIO(book)).sheetnames
        assert "Framework" in sheets, "the workbook does not name the reading"
        assert "Cognitive mix" in sheets
        assert "Mastery" not in sheets, "a cognitive export claims a KKM verdict"

        kkm_book = analysis_report.analysis_xlsx(
            analysis, self.EXAM, lang="en", framework=af.resolve("mastery"))
        kkm_sheets = load_workbook(BytesIO(kkm_book)).sheetnames
        assert "Mastery" in kkm_sheets
        assert "Cognitive mix" not in kkm_sheets

    def test_the_pdf_prints_the_framework_and_what_it_cannot_do(self, analysis):
        from app.services import analysis_report
        for key in ("ctt", "rasch", "cognitive", "mastery"):
            framework = af.resolve(key)
            text = pdf_text(analysis_report.analysis_pdf(
                analysis, self.EXAM, lang="en", framework=framework))
            assert framework.name["en"] in text, key
            # The honest half is in the filed document, not only in the browser.
            assert framework.not_for["en"][:40] in text.replace("\n", " "), key

    def test_the_pdf_prints_only_the_blocks_the_framework_publishes(self, analysis):
        from app.services import analysis_report
        labels = analysis_report.labels("en")
        cognitive = pdf_text(analysis_report.analysis_pdf(
            analysis, self.EXAM, lang="en", framework=af.resolve("cognitive")))
        classical = pdf_text(analysis_report.analysis_pdf(
            analysis, self.EXAM, lang="en", framework=af.resolve("ctt")))
        assert labels["cognitive"] in cognitive
        assert labels["cognitive"] not in classical
        assert labels["separation_block"] not in classical
        assert labels["separation_block"] in pdf_text(
            analysis_report.analysis_pdf(analysis, self.EXAM, lang="en",
                                         framework=af.resolve("rasch")))

    def test_every_route_that_builds_a_document_hands_over_the_framework(self):
        """Six call sites: three downloads for the teacher, three for a share link.

        They are two routes because they redact differently, and the one that
        drifts is the one nobody is looking at — a share link whose file names a
        different framework from the page it came from. Read per *call*, not by
        counting a substring: `framework=framework` also appears in the page's own
        `render_template`, so a count can stay green while a document loses it.
        """
        builders = ("analysis_report.analysis_csv(", "analysis_report.analysis_xlsx(",
                    "analysis_report.analysis_pdf(")
        for name in ("teacher.py", "public.py"):
            source = (ROOT / "app" / "routes" / name).read_text(encoding="utf-8")
            for builder in builders:
                for args in _call_arguments(source, builder):
                    assert "framework" in args, \
                        f"{name}: {builder} does not name the framework"
            assert "analysis_frameworks.resolve(request.args.get(\"framework\"))" in source

    def test_a_note_is_printed_only_where_its_subject_is(self, analysis):
        """`levels_partly_set` explains the HOTS shares; it is noise on a report
        that draws no shares, and noise in the notes is how a reader learns to skip
        them."""
        assert "levels_partly_set" in analysis.notes, "the fixture stopped testing it"
        ctt = af.notes_for(analysis.notes, af.resolve("ctt"))
        assert "levels_partly_set" not in ctt
        assert "levels_partly_set" in af.notes_for(analysis.notes, af.resolve("cognitive"))
        # Nothing is dropped that no framework owns: an unclassified note falls to
        # `summary`, which every framework shows.
        assert "dichotomous" in ctt

    def test_the_page_and_the_files_show_the_same_notes(self, analysis):
        from app.routes.teacher import _chart_payload
        chart = _chart_payload(analysis, self.EXAM, framework=af.resolve("ctt"))
        assert "levels_partly_set" not in chart["notes"]
        cognitive = _chart_payload(analysis, self.EXAM,
                                   framework=af.resolve("cognitive"))
        assert "levels_partly_set" in cognitive["notes"]


class TestTheBuilderRenders:
    """The page a teacher opens every day, rendered rather than read.

    A `{{ cognitive_levels() }}` that names a global the app never registered is
    not a wrong label — it is a 500 on the exam builder, on every subject, for
    every teacher. Reading the template as text cannot see that, so this renders
    it the way the route does and checks the six levels actually arrived.
    """

    @pytest.fixture(scope="module")
    def rendered(self):
        import contextlib

        from flask import g

        from app import create_app

        app = create_app("app.config.TestingConfig")
        with app.test_request_context("/teacher/exams/new"):
            g.user_id = "tea-1"
            g.user_name = "Guru Uji"
            g.user_email = "guru@example.test"
            g.user_role = "guru"
            g.tz_offset = 7
            g.show = {}
            with contextlib.suppress(Exception):
                g.user_school_id = "sch-1"
            return app.jinja_env.get_template("teacher/exam_form.html").render(
                exam=None, subjects=[], classes=[])

    def test_the_builder_renders_at_all(self, rendered):
        assert "SG_LEVELS" in rendered

    def test_all_six_levels_reach_the_page(self, rendered):
        """One entry per level, and each names the language pair, not one word."""
        listed = rendered.split("const SG_LEVELS = ", 1)[1].split(";", 1)[0]
        for item in af.LEVELS:
            assert item.key in listed
            assert item.name["id"] in listed and item.name["en"] in listed
        assert listed.count('"key"') == 6

    def test_every_panel_has_a_marker_and_every_marker_is_a_panel(self):
        """Both directions of the promise.

        A framework that claims a panel with no `data-panel` block shows a blank
        page and no error; a block with no marker is a block no framework can ever
        hide — including the ones that must not show it. The template's markers
        come from `af.PANELS`, so the two are checked against each other here.
        """
        text = (ROOT / "app" / "templates" / "teacher" / "analysis.html") \
            .read_text(encoding="utf-8")
        markers = set(re.findall(r'data-panel="(\w+)"', text))
        assert markers == set(af.PANELS), {
            "promised but not rendered": sorted(set(af.PANELS) - markers),
            "rendered but not promised": sorted(markers - set(af.PANELS))}

    def test_the_select_offers_every_level_as_an_option(self, rendered):
        select = rendered.split('<select x-model="q.level"', 1)[1].split("</select>", 1)[0]
        assert 'x-for="lv in SG_LEVELS"' in select
        # A brand-new paper labels nothing until the teacher says so.
        assert '<option value=""' in select

    def test_the_page_can_post_the_labels_it_collected(self, rendered):
        """The list the select fills has to reach the field the form posts."""
        assert 'name="question_cognitive"' in rendered
        assert "getLevelsJson()" in rendered
